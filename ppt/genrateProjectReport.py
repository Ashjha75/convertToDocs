#!/usr/bin/env python3
"""
PowerPoint Generator - THEMED (uses user's extracted palette)
Updated: uses more exact colors from extraction, improved fonts/sizes/paddings
"""

import argparse
import json
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml import parse_xml
import os

# Slide dimensions
SLIDE_W = Inches(10)
SLIDE_H = Inches(5.625)

# Better default font for PPTX (widely available)
FONT_FAMILY = "Calibri"

# Helper: convert hex -> RGBColor
def hex_to_rgb_tuple(hex_str):
    h = hex_str.lstrip("#")
    return tuple(int(h[i:i+2], 16) for i in (0, 2, 4))

def hex_to_rgbcolor(hex_str):
    r, g, b = hex_to_rgb_tuple(hex_str)
    return RGBColor(r, g, b)

# ---- THEME / PALETTE (expanded using your extracted values) ----
THEME = {
    # core theme slots (from your extraction)
    "dk1": "#000000",
    "lt1": "#FFFFFF",
    "dk2": "#1F497D",
    "lt2": "#EEECE1",
    "accent1": "#4F81BD",
    "accent2": "#C0504D",
    "accent3": "#9BBB59",
    "accent4": "#8064A2",
    "accent5": "#4BACC6",
    "accent6": "#F79646",
    "hlink": "#0000FF",
    "folHlink": "#800080",
    # table-related & extras (from your slide extraction)
    "table_row_alt": "#F3F4F6",
    "table_border": "#D1D5E3",
    "muted_gray": "#505050",
    "dark_text": "#15181F",
    # Additional frequently found colors (kept if present)
    "gray_light": "#E8E9F3",
    "danger": "#EF4444",
    "success": "#16A34A"
}

# PALETTE using theme (RGBColor instances)
PALETTE = {
    "ink": hex_to_rgbcolor(THEME["dk1"]),
    "text_dark": hex_to_rgbcolor(THEME["dark_text"]),
    "muted": hex_to_rgbcolor(THEME["muted_gray"]),
    "bg": hex_to_rgbcolor(THEME["lt1"]),
    "card": hex_to_rgbcolor(THEME["lt2"]),
    "accent": hex_to_rgbcolor(THEME["accent1"]),
    "accent2": hex_to_rgbcolor(THEME["accent5"]),
    "table_even": hex_to_rgbcolor("#FFFFFF"),
    "table_odd": hex_to_rgbcolor(THEME["table_row_alt"]),
    "chip": hex_to_rgbcolor(THEME["gray_light"] if "gray_light" in THEME else THEME["accent5"]),
    "table_border": THEME["table_border"],  # keep hex for XML insertion
    "success": hex_to_rgbcolor(THEME.get("success", "#16A34A")),
    "danger": hex_to_rgbcolor(THEME.get("danger", "#EF4444")),
}

# ---- Layout & sizing tweaks ----
# Slightly reduced title font, more space, smaller headings/body to avoid overflow
TITLE_FONT_SIZE = 36   # previously 44
SLIDE_TITLE_TOP = Inches(0.35)
DESC_BOX_HEIGHT = Inches(0.45)
PILL_FONT_SIZE = 10
HEADER_FONT_SIZE = 11
BODY_FONT_SIZE = 10
NOTE_FONT_SIZE = 9

def style_paragraph(paragraph, *, text="", size=10, bold=False, color=None,
                    align=None, line_spacing=1.15):
    paragraph.text = text
    paragraph.font.name = FONT_FAMILY
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = color or PALETTE["ink"]
    paragraph.line_spacing = line_spacing
    if align:
        paragraph.alignment = align

def set_slide_background_to_theme(slide):
    try:
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = PALETTE["bg"]
    except Exception:
        pass

def add_pill_text(slide, text, left, top, width, height):
    pill = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        left, top, width, height
    )
    pill.fill.solid()
    pill.fill.fore_color.rgb = PALETTE["card"]
    pill.line.fill.background()
    frame = pill.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.18)
    frame.margin_right = Inches(0.18)
    frame.margin_top = Inches(0.08)
    frame.margin_bottom = Inches(0.08)
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    style_paragraph(frame.paragraphs[0], text=text, size=PILL_FONT_SIZE, color=PALETTE["text_dark"], align=PP_ALIGN.CENTER)
    return pill

def add_table_borders(table):
    """Add borders to all table cells using theme border color"""
    val = PALETTE["table_border"].lstrip("#")
    for row in table.rows:
        for cell in row.cells:
            tc = cell._tc
            tcPr = tc.get_or_add_tcPr()
            # remove existing explicit lines if any? we append so keep previous behaviour
            for border in ['lnL', 'lnR', 'lnT', 'lnB']:
                ln = parse_xml(
                    f'<a:{border} xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" w="12700">'
                    f'<a:solidFill><a:srgbClr val="{val}"/></a:solidFill>'
                    f'</a:{border}>'
                )
                tcPr.append(ln)

def create_title_slide(prs, slide_data):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background_to_theme(slide)

    left_logo = slide_data.get("left_logo", "")
    right_logo = slide_data.get("right_logo", "")

    if left_logo and os.path.exists(left_logo):
        try:
            slide.shapes.add_picture(left_logo, Inches(0.5), Inches(0.35), width=Inches(1.8))
        except Exception:
            pass

    if right_logo and os.path.exists(right_logo):
        try:
            slide.shapes.add_picture(right_logo, Inches(7.5), Inches(0.35), width=Inches(1.8))
        except Exception:
            pass

    title_text = slide_data.get("title", "Project Report - Biller Management Portal")
    # widen height for long titles to avoid clipping
    title_box = slide.shapes.add_textbox(Inches(0.65), Inches(1.2), Inches(8.7), Inches(1.9))
    frame = title_box.text_frame
    frame.word_wrap = True
    style_paragraph(frame.paragraphs[0], text=title_text, size=TITLE_FONT_SIZE, bold=True, color=PALETTE["text_dark"], line_spacing=1.08)

    subtitle_text = slide_data.get("subtitle", "")
    if subtitle_text:
        add_pill_text(slide, subtitle_text, Inches(0.65), Inches(3.15), Inches(8.7), Inches(0.8))

def create_timeline_slide(prs, slide_data):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background_to_theme(slide)

    # Title
    title_box = slide.shapes.add_textbox(Inches(0.5), SLIDE_TITLE_TOP, Inches(9.0), Inches(0.6))
    p = title_box.text_frame.paragraphs[0]
    p.text = slide_data.get("title", "Project Timeline & Current Status")
    p.font.name = FONT_FAMILY
    p.font.size = Pt(TITLE_FONT_SIZE)
    p.font.bold = True
    p.font.color.rgb = PALETTE["text_dark"]

    # Description box - card color
    desc_text = slide_data.get("description", "")
    if desc_text:
        desc_shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(0.5), Inches(0.9),
            Inches(9.0), DESC_BOX_HEIGHT
        )
        desc_shape.fill.solid()
        desc_shape.fill.fore_color.rgb = PALETTE["card"]
        desc_shape.line.fill.background()

        tf = desc_shape.text_frame
        tf.clear()
        tf.word_wrap = True
        tf.margin_left = Inches(0.18)
        tf.margin_right = Inches(0.18)
        tf.margin_top = Inches(0.08)
        tf.margin_bottom = Inches(0.08)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE

        p = tf.paragraphs[0]
        p.text = desc_text
        p.font.name = FONT_FAMILY
        p.font.size = Pt(BODY_FONT_SIZE)
        p.font.color.rgb = PALETTE["text_dark"]
        p.alignment = PP_ALIGN.CENTER

    rows_data = slide_data.get("rows", [])
    if not rows_data:
        return

    num_rows = len(rows_data) + 1
    num_cols = 4
    left = Inches(0.5)
    top = Inches(1.6)
    width = Inches(9.0)
    height = Inches(3.5)

    table_shape = slide.shapes.add_table(num_rows, num_cols, left, top, width, height)
    table = table_shape.table

    # column widths tuned to reduce wrapping
    table.columns[0].width = Inches(2.1)  # Phase
    table.columns[1].width = Inches(2.4)  # Timeline
    table.columns[2].width = Inches(1.6)  # Status
    table.columns[3].width = Inches(2.9)  # Notes

    # Header row (accent color)
    headers = ["Phase", "Timeline", "Status", "Notes"]
    for col_idx, header_text in enumerate(headers):
        cell = table.cell(0, col_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = PALETTE["accent"]
        tf = cell.text_frame
        tf.clear()
        tf.margin_left = Inches(0.08)
        tf.margin_top = Inches(0.06)
        p = tf.paragraphs[0]
        p.text = header_text
        p.font.name = FONT_FAMILY
        p.font.size = Pt(HEADER_FONT_SIZE)
        p.font.bold = True
        p.font.color.rgb = PALETTE["bg"]

    # Data rows
    for row_idx, row_data in enumerate(rows_data, start=1):
        phase = row_data.get("phase", "")
        timeline = row_data.get("timeline", "")
        status = row_data.get("status", "")
        notes = row_data.get("notes", "")

        row_fill = PALETTE["table_odd"] if (row_idx % 2 == 1) else PALETTE["table_even"]
        for col_idx in range(num_cols):
            cell = table.cell(row_idx, col_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = row_fill

        # Phase
        cell = table.cell(row_idx, 0)
        tf = cell.text_frame
        tf.clear()
        tf.word_wrap = True
        tf.margin_left = Inches(0.08)
        tf.margin_top = Inches(0.06)
        p = tf.paragraphs[0]
        p.text = phase
        p.font.name = FONT_FAMILY
        p.font.size = Pt(BODY_FONT_SIZE)
        p.font.color.rgb = PALETTE["text_dark"]

        # Timeline
        cell = table.cell(row_idx, 1)
        tf = cell.text_frame
        tf.clear()
        tf.word_wrap = True
        tf.margin_left = Inches(0.08)
        tf.margin_top = Inches(0.06)
        p = tf.paragraphs[0]
        p.text = timeline
        p.font.name = FONT_FAMILY
        p.font.size = Pt(BODY_FONT_SIZE)
        p.font.color.rgb = PALETTE["text_dark"]

        # Status
        cell = table.cell(row_idx, 2)
        tf = cell.text_frame
        tf.clear()
        tf.word_wrap = True
        tf.margin_left = Inches(0.08)
        tf.margin_top = Inches(0.06)
        p = tf.paragraphs[0]
        status_lower = (status or "").lower()
        if "completed" in status_lower and "partially" not in status_lower:
            p.text = "✓ " + status
            status_color = PALETTE["success"]
        elif "partially" in status_lower or "review" in status_lower:
            p.text = "✓ " + status
            status_color = PALETTE["success"]
        elif "progress" in status_lower:
            p.text = "◐ " + status
            status_color = PALETTE["accent"]
        elif "not started" in status_lower:
            p.text = "✗ " + status
            status_color = PALETTE["danger"]
        else:
            p.text = status
            status_color = PALETTE["text_dark"]
        p.font.name = FONT_FAMILY
        p.font.size = Pt(BODY_FONT_SIZE)
        p.font.color.rgb = status_color

        # Notes
        cell = table.cell(row_idx, 3)
        tf = cell.text_frame
        tf.clear()
        tf.word_wrap = True
        tf.margin_left = Inches(0.08)
        tf.margin_top = Inches(0.06)
        p = tf.paragraphs[0]
        p.text = notes
        p.font.name = FONT_FAMILY
        p.font.size = Pt(NOTE_FONT_SIZE)
        p.font.color.rgb = PALETTE["muted"]

    add_table_borders(table)

def create_three_column_slide(prs, slide_data):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background_to_theme(slide)

    title_box = slide.shapes.add_textbox(Inches(0.5), SLIDE_TITLE_TOP, Inches(9.0), Inches(0.6))
    p = title_box.text_frame.paragraphs[0]
    p.text = slide_data.get("title", "Our Solution: A Unified Three Pillar Platform")
    p.font.name = FONT_FAMILY
    p.font.size = Pt(TITLE_FONT_SIZE)
    p.font.bold = True
    p.font.color.rgb = PALETTE["text_dark"]

    desc_text = slide_data.get("description", "")
    if desc_text:
        desc_shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(0.5), Inches(0.9),
            Inches(9.0), DESC_BOX_HEIGHT
        )
        desc_shape.fill.solid()
        desc_shape.fill.fore_color.rgb = PALETTE["card"]
        desc_shape.line.fill.background()

        tf = desc_shape.text_frame
        tf.clear()
        tf.word_wrap = True
        tf.margin_left = Inches(0.18)
        tf.margin_right = Inches(0.18)
        tf.margin_top = Inches(0.08)
        tf.margin_bottom = Inches(0.08)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE

        p = tf.paragraphs[0]
        p.text = desc_text
        p.font.name = FONT_FAMILY
        p.font.size = Pt(BODY_FONT_SIZE)
        p.font.color.rgb = PALETTE["text_dark"]
        p.alignment = PP_ALIGN.CENTER

    columns = slide_data.get("columns", [])
    if not columns or len(columns) != 3:
        return

    col_width = Inches(2.8)
    col_gap = Inches(0.28)
    start_x = Inches(0.5)
    start_y = Inches(1.6)

    for i, col_data in enumerate(columns):
        x_pos = start_x + i * (col_width + col_gap)
        heading_box = slide.shapes.add_textbox(x_pos, start_y, col_width, Inches(0.4))
        tf = heading_box.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = col_data.get("heading", "")
        p.font.name = FONT_FAMILY
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = PALETTE["text_dark"]

        bullets = col_data.get("bullets", [])
        bullets_box = slide.shapes.add_textbox(x_pos, start_y + Inches(0.5), col_width, Inches(3.3))
        tf = bullets_box.text_frame
        tf.word_wrap = True
        tf.clear()

        for idx, bullet_text in enumerate(bullets):
            if idx == 0:
                p = tf.paragraphs[0]
            else:
                p = tf.add_paragraph()
            p.text = bullet_text
            p.font.name = FONT_FAMILY
            p.font.size = Pt(NOTE_FONT_SIZE)
            p.font.color.rgb = PALETTE["text_dark"]
            p.space_after = Pt(8)

def create_table_slide(prs, slide_data):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background_to_theme(slide)

    title_box = slide.shapes.add_textbox(Inches(0.5), SLIDE_TITLE_TOP, Inches(9.0), Inches(0.6))
    p = title_box.text_frame.paragraphs[0]
    p.text = slide_data.get("title", "Modules & Deep Dive On Features")
    p.font.name = FONT_FAMILY
    p.font.size = Pt(TITLE_FONT_SIZE)
    p.font.bold = True
    p.font.color.rgb = PALETTE["text_dark"]

    desc_text = slide_data.get("description", "")
    if desc_text:
        desc_shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(0.5), Inches(0.9),
            Inches(9.0), DESC_BOX_HEIGHT
        )
        desc_shape.fill.solid()
        desc_shape.fill.fore_color.rgb = PALETTE["card"]
        desc_shape.line.fill.background()

        tf = desc_shape.text_frame
        tf.clear()
        tf.word_wrap = True
        tf.margin_left = Inches(0.18)
        tf.margin_right = Inches(0.18)
        tf.margin_top = Inches(0.06)
        tf.margin_bottom = Inches(0.06)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE

        p = tf.paragraphs[0]
        p.text = desc_text
        p.font.name = FONT_FAMILY
        p.font.size = Pt(BODY_FONT_SIZE)
        p.font.color.rgb = PALETTE["text_dark"]
        p.alignment = PP_ALIGN.CENTER

    columns = slide_data.get("columns", [])
    rows_data = slide_data.get("rows", [])
    if not columns or not rows_data:
        return

    num_rows = len(rows_data) + 1
    num_cols = len(columns)
    left = Inches(0.5)
    top = Inches(1.5)
    width = Inches(9.0)
    height = Inches(3.7)

    table_shape = slide.shapes.add_table(num_rows, num_cols, left, top, width, height)
    table = table_shape.table

    # sensible default widths (adjustable)
    if num_cols >= 3:
        table.columns[0].width = Inches(1.8)
        table.columns[1].width = Inches(5.4)
        table.columns[2].width = Inches(1.8)

    # header
    for col_idx, header_text in enumerate(columns):
        cell = table.cell(0, col_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = PALETTE["accent"]
        tf = cell.text_frame
        tf.clear()
        tf.margin_left = Inches(0.08)
        tf.margin_top = Inches(0.06)
        p = tf.paragraphs[0]
        p.text = header_text
        p.font.name = FONT_FAMILY
        p.font.size = Pt(HEADER_FONT_SIZE)
        p.font.bold = True
        p.font.color.rgb = PALETTE["bg"]

    # data rows
    for row_idx, row in enumerate(rows_data, start=1):
        row_fill = PALETTE["table_odd"] if (row_idx % 2 == 1) else PALETTE["table_even"]
        for col_idx in range(num_cols):
            cell = table.cell(row_idx, col_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = row_fill
            tf = cell.text_frame
            tf.clear()
            tf.word_wrap = True
            tf.margin_left = Inches(0.08)
            tf.margin_top = Inches(0.06)
            tf.margin_right = Inches(0.08)

            value = row[col_idx] if isinstance(row, list) else row.get(columns[col_idx], "")
            p = tf.paragraphs[0]
            p.text = str(value)
            p.font.name = FONT_FAMILY
            p.font.size = Pt(NOTE_FONT_SIZE)
            p.font.color.rgb = PALETTE["text_dark"]

    add_table_borders(table)

def generate_presentation(data_path, output_path):
    with open(data_path, 'r', encoding='utf-8') as f:
        data = json.load(f)
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    for slide_data in data.get("slides", []):
        slide_type = slide_data.get("type", "")
        if slide_type == "title":
            create_title_slide(prs, slide_data)
        elif slide_type == "timeline":
            create_timeline_slide(prs, slide_data)
        elif slide_type == "three_column":
            create_three_column_slide(prs, slide_data)
        elif slide_type == "table":
            create_table_slide(prs, slide_data)
    prs.save(output_path)
    print(f"✓ Presentation created successfully!")
    print(f"  File: {output_path}")
    print(f"  Slides: {len(prs.slides)}")

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Generate PowerPoint from JSON (themed)")
    parser.add_argument("--data", default="data.json", help="Input JSON file")
    parser.add_argument("--out", default="Final_Report_Themed_Updated.pptx", help="Output PPTX file")
    args = parser.parse_args()

    if not os.path.exists(args.data):
        print(f"ERROR: File not found: {args.data}")
        exit(1)

    generate_presentation(args.data, args.out)
