#!/usr/bin/env python3
# generate_biller_doc.py
import json
import os
import zipfile

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

# ---------- CONFIG ----------
INPUT_JSON = "rms-deposit.json"  # path to your JSON file
OUTPUT_DOCX = "Mpos-Biller-Portal-Test_Cases.docx"
OUTPUT_ZIP = "Mpos-Biller-Portal-Test_Cases.zip"
TITLE_TEXT = "Mpos-Biller Portal Test Cases"
FONT_NAME = "Calibri"
FONT_SIZE_PT = 11
PASS_FAIL_TEXT = "PASS\nFAIL\n\nCOMMENTS:"
# ----------------------------


def load_json(path: str):
    with open(path, "r", encoding="utf-8") as f:
        content = f.read().strip()
    try:
        return json.loads(content)
    except json.JSONDecodeError as err:
        try:
            return json.loads("[" + content + "]")
        except Exception as wrap_err:
            raise err from wrap_err


def set_run_font(run, bold: bool = True) -> None:
    """Apply consistent font styling to a run."""
    run.font.name = FONT_NAME
    run.font.size = Pt(FONT_SIZE_PT)
    run.font.bold = bold


def set_cell_font(cell) -> None:
    """Apply consistent font styling to all runs in a cell."""
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            set_run_font(run)


def main() -> None:
    """Generate Word document with test cases from JSON data."""
    data = load_json(INPUT_JSON)
    if not isinstance(data, list):
        raise SystemExit("JSON root must be an array of test case objects.")

    doc = Document()

    section = doc.sections[-1]
    section.orientation = WD_ORIENT.LANDSCAPE
    # Set to Legal size (21.59 cm x 35.56 cm)
    section.page_width = Inches(14.0)  # 35.56 cm ≈ 14 inches
    section.page_height = Inches(8.5)  # 21.59 cm ≈ 8.5 inches
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)

    style = doc.styles["Normal"]
    style.font.name = FONT_NAME
    style.font.size = Pt(FONT_SIZE_PT)
    style.font.bold = True

    title = doc.add_heading(TITLE_TEXT, level=1)
    for run in title.runs:
        set_run_font(run)
        run.font.size = Pt(14)

    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    header_labels = [
        "Test Case ID",
        "Test Case",
        "Test Procedures / Steps",
        "Expected Outcome",
        "Pass/Fail",
    ]
    header_cells = table.rows[0].cells
    for idx, text in enumerate(header_labels):
        cell = header_cells[idx]
        cell.text = text
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.name = FONT_NAME
                run.font.size = Pt(12)
                run.font.bold = True
                run.font.color.rgb = RGBColor(153, 184, 235)

    row_id = 1
    for item in data:
        cells = table.add_row().cells
        cells[0].text = str(row_id)
        cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_cell_font(cells[0])

        cells[1].text = item.get("Test Case", "")
        set_cell_font(cells[1])

        procedures = cells[2]
        procedures.text = ""
        steps = item.get("Test Procedures/Steps", [])
        if isinstance(steps, list) and steps:
            for step in steps:
                cleaned = step.lstrip("0123456789. ").strip()
                paragraph = procedures.add_paragraph(style="List Bullet")
                run = paragraph.add_run(cleaned)
                set_run_font(run)
        else:
            procedures.text = str(steps)
            for paragraph in procedures.paragraphs:
                for run in paragraph.runs:
                    set_run_font(run)

        cells[3].text = item.get("Expected Outcome", "")
        set_cell_font(cells[3])

        cells[4].text = PASS_FAIL_TEXT
        set_cell_font(cells[4])

        row_id += 1

    doc.save(OUTPUT_DOCX)

    with zipfile.ZipFile(OUTPUT_ZIP, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.write(OUTPUT_DOCX, arcname=os.path.basename(OUTPUT_DOCX))

    print("Done.")
    print("DOCX:", os.path.abspath(OUTPUT_DOCX))
    print("ZIP:", os.path.abspath(OUTPUT_ZIP))


if __name__ == "__main__":
    main()
